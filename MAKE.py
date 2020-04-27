from openpyxl import load_workbook

wb = load_workbook('movie_optimization.xlsx')
print(wb.sheetnames())
